"""v0.8 regression (code audit H1, H2, M9): the honesty audit reads every text channel and fails closed.
Builds from the shipped jordan_example workspace, so it needs the engine venv (python-docx).
Run: <jobkit python> -m unittest system/engine/tools/tests/test_audit_channels.py"""
import glob, os, shutil, subprocess, sys, tempfile, unittest

TOOLS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = os.path.abspath(os.path.join(TOOLS, "..", "..", ".."))
WS = os.path.join(ROOT, "system", "examples", "jordan_example")
SRC = sorted(glob.glob(os.path.join(WS, "applications", "*", "Resume_Jordan_Example.docx")))


def audit(folder):
    out = subprocess.run([sys.executable, os.path.join(TOOLS, "audit_claims.py"), folder],
                         env=dict(os.environ, JOBKIT_WORKSPACE=WS), capture_output=True, text=True)
    return out.stdout + out.stderr


@unittest.skipUnless(SRC, "jordan_example has no built resume (run its build_test_example.py first)")
class TestAuditChannels(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        from docx import Document
        self.Document = Document

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _copy(self, name):
        d = os.path.join(self.tmp, "f"); os.makedirs(d, exist_ok=True)
        p = os.path.join(d, name); shutil.copy(SRC[0], p); return d, p

    def test_clean_resume_still_passes(self):
        d, _ = self._copy("Resume_Jordan_Example.docx")
        self.assertIn("PASS", audit(d)); self.assertTrue(os.path.exists(os.path.join(d, "AUDIT_PASS.json")))

    def test_forbidden_claim_in_table_fails(self):
        d, p = self._copy("Resume_Jordan_Example.docx"); doc = self.Document(p)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "Nurse practitioner; reduced CLABSI 60%"
        doc.save(p)
        out = audit(d)
        self.assertIn("FAIL", out); self.assertIn("forbidden", out); self.assertFalse(os.path.exists(os.path.join(d, "AUDIT_PASS.json")))

    def test_claim_in_footer_fails(self):
        d, p = self._copy("Resume_Jordan_Example.docx"); doc = self.Document(p)
        doc.sections[0].footer.paragraphs[0].text = "Flight nurse, intubation-certified, 15 years in trauma"
        doc.save(p)
        out = audit(d)
        self.assertIn("FAIL", out); self.assertIn("outside the page body", out)

    def test_unapproved_line_in_header_fails(self):
        d, p = self._copy("Resume_Jordan_Example.docx"); doc = self.Document(p)
        doc.sections[0].header.paragraphs[0].text = "Charge nurse manager of a 40-bed ICU"   # no forbidden term; still unapproved
        doc.save(p)
        self.assertIn("FAIL", audit(d))

    def test_renamed_file_is_still_a_resume(self):
        d, p = self._copy("Jordan_Resume.docx"); doc = self.Document(p)
        b = [x for x in doc.paragraphs if x.style.name.startswith("List")][0]
        b.runs[0].text = "Directed the hospital's entire critical-care division and built its ECMO program from scratch."
        for r in b.runs[1:]: r.text = ""
        doc.save(p)
        out = audit(d)
        self.assertIn("FAIL", out); self.assertIn("traceability", out)

    def test_stamped_letter_is_checked_as_a_letter(self):
        d = os.path.join(self.tmp, "f"); os.makedirs(d, exist_ok=True)
        p = os.path.join(d, "CoverLetter_Jordan_Example.docx")
        doc = self.Document(); doc.add_paragraph("Dear hiring manager, I am a medical ICU nurse and would like to apply.")
        sys.path.insert(0, TOOLS); import docxtext
        docxtext.stamp(doc, "letter"); doc.save(p)
        self.assertIn("PASS", audit(d))

    def test_lock_file_and_backup_are_ignored(self):
        d, p = self._copy("Resume_Jordan_Example.docx")
        shutil.copy(p, os.path.join(d, "~$Resume_Jordan_Example.docx"))
        bak = os.path.join(d, "old.bak-2026-01-01"); os.makedirs(bak); shutil.copy(p, os.path.join(bak, "Resume_Jordan_Example.docx"))
        with open(os.path.join(d, "~$Resume_Jordan_Example.docx"), "wb") as f: f.write(b"not a docx")
        out = audit(d)
        self.assertNotIn("~$", out); self.assertNotIn("bak-", out); self.assertIn("1/1 clean", out)

    def test_unreadable_file_is_a_fail_line_not_a_crash(self):
        d, p = self._copy("Resume_Jordan_Example.docx")
        with open(p, "wb") as f: f.write(b"garbage")
        out = audit(d)
        self.assertIn("FAIL", out); self.assertIn("unreadable", out); self.assertNotIn("Traceback", out)


if __name__ == "__main__":
    unittest.main()


@unittest.skipUnless(SRC, "jordan_example has no built resume")
class TestStructuralExemptions(unittest.TestCase):
    """v0.8 (code audit H6): a claim cannot hide in an ALL-CAPS 'heading', a line next to the email, a label, or a summary prefix."""
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        from docx import Document
        self.Document = Document
        self.d = os.path.join(self.tmp, "f"); os.makedirs(self.d)
        self.p = os.path.join(self.d, "Resume_Jordan_Example.docx"); shutil.copy(SRC[0], self.p)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _insert(self, text, after=0):
        doc = self.Document(self.p)
        anchor = doc.paragraphs[after]
        new = anchor.insert_paragraph_before(text) if after == 0 else None
        if new is None:
            from copy import deepcopy
            el = deepcopy(anchor._p); anchor._p.addnext(el)
            from docx.text.paragraph import Paragraph
            q = Paragraph(el, anchor._parent)
            for r in q.runs[1:]: r.text = ""
            q.runs[0].text = text
        doc.save(self.p)

    def test_all_caps_claim_is_not_a_heading(self):
        self._insert("RN, BSN, 12 YEARS ICU CHARGE", after=3)
        self.assertIn("FAIL", audit(self.d))

    def test_claim_next_to_email_is_not_a_contact_line(self):
        self._insert("Led 40 nurses as unit director | jordan.example@example.com", after=3)
        self.assertIn("FAIL", audit(self.d))

    def test_arbitrary_label_before_approved_text_fails(self):
        # an approved skills text behind a label that is not an approved label
        doc = self.Document(self.p)
        skills = [q for q in doc.paragraphs if q.text.strip().lower().startswith("licenses & certifications")][0]
        skills.runs[0].text = "Charge Nurse Manager, 40-bed ICU : " + skills.text.split(":", 1)[1].strip()
        for r in skills.runs[1:]: r.text = ""
        doc.save(self.p)
        self.assertIn("FAIL", audit(self.d))

    def test_summary_prefix_with_a_number_fails(self):
        doc = self.Document(self.p)
        summ = [q for q in doc.paragraphs if "registered nurse" in q.text.lower() and len(q.text) > 120][0]
        body = summ.text.split(": ", 1)[1]
        summ.runs[0].text = "Nurse leader who cut CLABSI 60% in one year: " + body
        for r in summ.runs[1:]: r.text = ""
        doc.save(self.p)
        self.assertIn("FAIL", audit(self.d))

    def test_plain_role_noun_prefix_still_passes(self):
        self.assertIn("PASS", audit(self.d))   # the shipped resume opens "Medical ICU registered nurse candidate: <summary>"


class TestLoaderReservedNames(unittest.TestCase):
    """v0.8 (code audit H7): a bullet named like an engine constant is refused with a clear error, never a crash or a silent swap."""
    def test_reserved_name_is_refused(self):
        ws = tempfile.mkdtemp()
        try:
            os.makedirs(os.path.join(ws, "profile"))
            with open(os.path.join(ws, "settings.yaml"), "w") as f:
                f.write('identity: {name: "T P", first_name: T, email: t@x.com, phone: "1", city_state: "A, PA"}\nwork_authorization: {resume_line: "U.S. Citizen"}\nlanes: []\n')
            with open(os.path.join(ws, "profile", "approved_claims.yaml"), "w") as f:
                f.write('experience:\n  - key: SUMMARIES\n    org: X\n    dates: "2020 - 2021"\n    role: R\n    loc: L\n    bullets: {EXPERIENCE: "Did a thing."}\n')
            # v1.0 (audit M26): the refusal fires on the first LOAD (load(), a name import, Resume(), the audit), no longer on
            # `import approved_claims` itself, which reads nothing. Same wording, same non-zero exit.
            out = subprocess.run([sys.executable, "-c", f"import sys; sys.path.insert(0, {TOOLS!r}); import approved_claims; approved_claims.load()"],
                                 env=dict(os.environ, JOBKIT_WORKSPACE=ws), capture_output=True, text=True)
            self.assertNotEqual(out.returncode, 0)
            self.assertIn("reserved", out.stderr); self.assertIn("SUMMARIES", out.stderr)
        finally:
            shutil.rmtree(ws, ignore_errors=True)


class TestLazyLoader(unittest.TestCase):
    """v1.0 (audit M26): importing approved_claims / audit_claims reads no file; claims_loader.load() is explicit, cached and
    shared; a missing YAML is one sentence; a reserved name is refused by every entry point that builds or audits."""
    SETTINGS = 'identity: {name: "T P", first_name: T, email: t@x.com, phone: "1", city_state: "A, PA"}\nwork_authorization: {resume_line: "U.S. Citizen"}\nlanes: []\n'
    RESERVED = 'experience:\n  - key: SUMMARIES\n    org: X\n    dates: "2020 - 2021"\n    role: R\n    loc: L\n    bullets: {EXPERIENCE: "Did a thing."}\n'

    def setUp(self):
        self.ws = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.ws, ignore_errors=True)

    def _run(self, code):
        return subprocess.run([sys.executable, "-c", f"import sys; sys.path.insert(0, {TOOLS!r})\n" + code],
                              env=dict(os.environ, JOBKIT_WORKSPACE=self.ws), capture_output=True, text=True)

    def _reserved_ws(self):
        os.makedirs(os.path.join(self.ws, "profile"))
        with open(os.path.join(self.ws, "settings.yaml"), "w") as f: f.write(self.SETTINGS)
        with open(os.path.join(self.ws, "profile", "approved_claims.yaml"), "w") as f: f.write(self.RESERVED)

    def test_import_in_empty_workspace_is_silent(self):
        out = self._run("import approved_claims, audit_claims, resume_builder, consistency, keywords; print('imported')")
        self.assertEqual(out.returncode, 0, out.stderr)
        self.assertEqual(out.stdout.strip(), "imported"); self.assertEqual(out.stderr, "")

    def test_import_does_not_validate_the_yaml(self):
        self._reserved_ws()   # a YAML the loader refuses: import must still succeed, so nothing was read at import
        out = self._run("import approved_claims, audit_claims; print('imported')")
        self.assertEqual(out.returncode, 0, out.stderr); self.assertEqual(out.stdout.strip(), "imported")

    def test_loader_in_empty_workspace_gives_a_sentence(self):
        out = self._run("import approved_claims as a, claims_loader as cl\n"
                        "c = a.load(); assert not c.loaded() and not a.loaded(); assert a.load() is c and cl.load() is c\n"
                        "print(c.missing_sentence())\n"
                        "try:\n    from approved_claims import ERMC\nexcept ImportError as e:\n    print('IMPORTERR', e)\n")
        self.assertEqual(out.returncode, 0, out.stderr)
        self.assertIn("approved_claims.yaml is missing", out.stdout); self.assertIn("run onboarding", out.stdout)
        self.assertIn("IMPORTERR approved_claims has no name 'ERMC'", out.stdout); self.assertNotIn("Traceback", out.stderr)

    def test_audit_cli_refuses_reserved_name_before_reading_any_docx(self):
        self._reserved_ws()
        out = subprocess.run([sys.executable, os.path.join(TOOLS, "audit_claims.py"), self.ws],
                             env=dict(os.environ, JOBKIT_WORKSPACE=self.ws), capture_output=True, text=True)
        self.assertEqual(out.returncode, 2, out.stderr)
        self.assertIn("reserved", out.stderr); self.assertIn("SUMMARIES", out.stderr); self.assertNotIn("Traceback", out.stderr)

    def test_builder_and_name_import_refuse_reserved_name(self):
        self._reserved_ws()
        for code in ("from resume_builder import Resume; Resume()", "from approved_claims import SUMMARIES"):
            out = self._run(code)
            self.assertNotEqual(out.returncode, 0, code); self.assertIn("reserved", out.stderr); self.assertIn("SUMMARIES", out.stderr)

    def test_facade_names_are_all_reserved(self):
        """A YAML name that shadowed a facade attribute would be served silently by the module, never by __getattr__."""
        sys.path.insert(0, TOOLS)
        import approved_claims, claims_loader
        public = {n for n in vars(approved_claims) if not n.startswith("__")}
        self.assertTrue(public <= claims_loader._ENGINE_NAMES, public - claims_loader._ENGINE_NAMES)
