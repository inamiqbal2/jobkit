#!/usr/bin/env python3
"""
Reusable one-page ATS resume builder (jobkit engine).
Identity, work-authorization line, and education come from My Job Search/settings.yaml (userconfig.py).
Nothing personal is hard-coded here.

Format is FIXED - do not change fonts/margins/spacing to fill space.
If a page is underfilled, ADD REAL SUBSTANCE (see system/engine/charters/build.md).
"""
import sys as _sys0
try: _sys0.stdout.reconfigure(encoding="utf-8", errors="replace")   # Windows: never die on a non-cp1252 character
except Exception: pass
from docx import Document
from docx.shared import Pt, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime, re, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from userconfig import get as _cfg, work_auth as _work_auth, build as _build_cfg
import approved_claims as _ac   # CV sections (optional `cv:` block)

NAME = _cfg("identity.name", "")
CITY = _cfg("identity.city_state", "")     # ATS location prefill needs a city/state; OFF the contact line by default
PHONE = _cfg("identity.phone", "")
EMAIL = _cfg("identity.email", "")
LINKS = list(_cfg("identity.links", []) or [])
EDUCATION = list(_cfg("education", []) or [])
WORK_AUTH_DEFAULT = _work_auth()["resume_line"] or None
MIN_JOBS = _build_cfg()["min_experience_entries"]
LOCATION_LINE_DEFAULT = _build_cfg()["location_line_default"]
RIGHT_TAB = 10080  # twips


class Resume:
    def __init__(self, location_line=None, work_auth=WORK_AUTH_DEFAULT, city=None):
        """
        location_line: contact-line text about location. Default = settings.yaml build.location_line_default
                       (e.g. "Open to relocation"). For commutable roles pass the metro name. Never the
                       word "local" (audit_claims bans it).
        work_auth: appended to the contact line. Default = settings.yaml work_authorization.resume_line.
                   Pass None to omit.
        city: ATS location prefill wants a city/state as the FIRST contact item (pass city=CITY).
              DEFAULT OFF because it may not fit on one line at the fixed 9.5 pt; measure with
              verify_resume.py before enabling.
        """
        if not NAME:
            raise RuntimeError("settings.yaml identity.name is empty - run the onboarding skill first")
        self.doc = self._new_doc()
        parts = [p for p in [PHONE, EMAIL] + LINKS if p]
        parts.append(location_line or LOCATION_LINE_DEFAULT)
        contact = "  |  ".join(parts)
        if city:
            contact = f"{city}  |  " + contact
        if work_auth:
            contact += f"  |  {work_auth}"
        self._name_block(contact)

    # ---------- setup ----------
    def _new_doc(self):
        doc = Document()
        sec = doc.sections[0]
        sec.top_margin = sec.bottom_margin = Pt(30)
        sec.left_margin = sec.right_margin = Pt(52)
        s = doc.styles["Normal"]
        s.font.name = "Calibri"
        s.font.size = Pt(10)
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_after = Pt(1)
        s.paragraph_format.space_before = Pt(0)
        s.paragraph_format.line_spacing = 1.0
        self._fix_bullet_glyph(doc)
        return doc

    @staticmethod
    def _fix_bullet_glyph(doc):
        """The default template's List Bullet level uses the Symbol-font private-use glyph U+F0B7,
        which PDF parsers emit as garbage. Swap the level text to Unicode U+2022 in the body font."""
        body_font = doc.styles["Normal"].font.name or "Calibri"
        numbering = doc.part.numbering_part.element
        for lvl in numbering.iter(qn("w:lvl")):
            fmt = lvl.find(qn("w:numFmt"))
            if fmt is None or fmt.get(qn("w:val")) != "bullet":
                continue
            txt = lvl.find(qn("w:lvlText"))
            if txt is None or txt.get(qn("w:val")) != "":
                continue
            txt.set(qn("w:val"), "•")
            rpr = lvl.find(qn("w:rPr"))
            if rpr is not None:
                for fonts in rpr.findall(qn("w:rFonts")):
                    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                        if fonts.get(qn(attr)) is not None:
                            fonts.set(qn(attr), body_font)
                    if fonts.get(qn("w:hint")) is not None:
                        del fonts.attrib[qn("w:hint")]

    def _name_block(self, contact):
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(NAME); r.bold = True; r.font.size = Pt(17)
        p2 = self.doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run(contact).font.size = Pt(9.5)

    def _heading(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(text.upper()); r.bold = True; r.font.size = Pt(11)
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "6")
        b.set(qn("w:space"), "1"); b.set(qn("w:color"), "000000")
        pbdr.append(b); pPr.append(pbdr)

    def _tabbed(self, space_before=0):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Twips(RIGHT_TAB), WD_TAB_ALIGNMENT.RIGHT)
        return p

    def _bullet(self, text):
        bp = self.doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(1)
        bp.paragraph_format.space_before = Pt(0)
        bp.add_run(text).font.size = Pt(10)

    # ---------- sections ----------
    def summary(self, text):
        self._heading("Summary")
        self.doc.add_paragraph(text)
        return self

    # ---- experience entries are COLLECTED here and rendered in reverse-chronological order ----
    _MONTHS = {m: i for i, m in enumerate(
        ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

    @classmethod
    def _date_key(cls, dates):
        """Sort key for a date string: ("Present" first), then start year, then start month.
        Accepts 'Jan 2025 - Present', 'Jun 2024 - Aug 2024', '2020 - 2021' (year-only = Jan)."""
        m = re.match(r"\s*(?:([A-Za-z]{3})[A-Za-z]*\.?\s+)?(\d{4})", dates or "")
        if not m:
            return (0, 0, 0)
        month = cls._MONTHS.get((m.group(1) or "jan").lower(), 1)
        present = 1 if re.search(r"\bpresent\b", dates, re.I) else 0
        return (present, int(m.group(2)), month)

    @staticmethod
    def _norm_dates(dates):
        """'Sept' -> 'Sep' (parsers prefix-match 3-letter months)."""
        return re.sub(r"\bSept\b\.?", "Sep", dates)

    def job(self, org, dates, role, loc, bullets):
        self._jobs = getattr(self, "_jobs", [])
        self._jobs.append((org, self._norm_dates(dates), role, loc, list(bullets)))
        self._job_count = len(self._jobs)
        return self

    def _render_job(self, org, dates, role, loc, bullets):
        p = self._tabbed(space_before=3)
        r = p.add_run(org); r.bold = True; r.font.size = Pt(10.5)
        p.add_run("\t" + dates).font.size = Pt(10)
        p2 = self.doc.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.tab_stops.add_tab_stop(Twips(RIGHT_TAB), WD_TAB_ALIGNMENT.RIGHT)
        ri = p2.add_run(role); ri.italic = True; ri.font.size = Pt(10)
        ri2 = p2.add_run("\t" + loc); ri2.italic = True; ri2.font.size = Pt(10)
        for b in bullets:
            self._bullet(b)

    def _flush_jobs(self):
        """Write the Experience section (heading + entries sorted by start date, descending)."""
        jobs = getattr(self, "_jobs", None)
        if not jobs:
            return
        ordered = sorted(jobs, key=lambda j: self._date_key(j[1]), reverse=True)  # stable
        if [j[0] for j in ordered] != [j[0] for j in jobs]:
            print("WARNING: experience entries re-ordered to reverse-chronological: "
                  + " > ".join(f"{j[0]} ({j[1]})" for j in ordered)
                  + "  [caller order was: " + " > ".join(j[0] for j in jobs) + "]")
        if not getattr(self, "_exp_heading_written", False):
            self._heading("Experience"); self._exp_heading_written = True
        for j in ordered:
            self._render_job(*j)
        self._jobs = []

    def experience_heading(self):
        self._heading("Experience"); self._exp_heading_written = True; return self

    def projects_heading(self, title="Projects"):
        self._flush_jobs()
        self._heading(title); return self

    def project(self, title, bullet):
        self._flush_jobs()
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(0)
        p.add_run(title).bold = True
        self._bullet(bullet)
        return self

    def education(self, entries=None):
        """Education entries from settings.yaml -> education: [{institution, dates, line}], newest first.
        Pass `entries` to override for one build (same shape)."""
        self._flush_jobs()
        self._heading("Education")
        for e in (entries or EDUCATION):
            p = self._tabbed()
            p.add_run(str(e.get("institution", ""))).bold = True
            p.add_run("\t" + str(e.get("dates", ""))).font.size = Pt(10)
            self.doc.add_paragraph(str(e.get("line", "")))
        return self

    def skills(self, lines):
        self._flush_jobs()
        self._heading("Skills")
        for label, val in lines:
            p = self.doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
            p.add_run(label + " : ").bold = True
            p.add_run(val)
        return self

    def _set_core_properties(self):
        """Strip the python-docx fingerprint (creator 'python-docx', created 2013-12-23)."""
        cp = self.doc.core_properties
        now = datetime.datetime.now().replace(microsecond=0)
        cp.author = NAME
        cp.title = f"{NAME} - Resume"
        cp.last_modified_by = NAME
        cp.created = now
        cp.modified = now
        cp.comments = ""
        cp.keywords = ""
        cp.subject = "jobkit:resume"   # v0.8: the audit reads this stamp to decide how strictly to check the file
        cp.category = ""
        cp.revision = 1

    def save(self, path):
        self._flush_jobs()
        if getattr(self, "_job_count", 0) < MIN_JOBS:
            print(f"WARNING: RULE VIOLATION (build.min_experience_entries={MIN_JOBS}): only "
                  f"{getattr(self,'_job_count',0)} work experiences. Cut projects/skills instead of a job entry.")
        self._set_core_properties()
        self.doc.save(path)
        print("wrote", path)
        return path


# ═══════════════════════════════════════════════════════════════════════════════════════════════════════════════
# ACADEMIC CV (2026-10-09): the multi-page format for research / PhD applications. Same type, same contact line, same
# audit; sections instead of a one-page squeeze. Page range is enforced by verify_resume.py (settings.yaml build.cv_min_pages
# .. cv_max_pages, default 2..3); EVERY line is traceable to approved_claims.yaml (audit_claims.py CV mode).
# settings.yaml build.format: "academic_cv" selects it; "resume" (default) keeps the one-page builder above untouched.
class AcademicCV(Resume):
    SECTION_TITLES = {"research": "Research Experience", "professional": "Professional Experience",
                      "field": "Field Experience & Expeditions", "teaching": "Teaching & Mentoring"}

    def __init__(self, location_line=None, work_auth=WORK_AUTH_DEFAULT, city=None, page_numbers=True):
        super().__init__(location_line=location_line, work_auth=work_auth, city=city)
        if page_numbers:
            self._page_footer()

    def _page_footer(self):
        """'Name - Curriculum Vitae - page N' in the footer (PAGE field; renders in Word and LibreOffice)."""
        from docx.oxml import OxmlElement
        foot = self.doc.sections[0].footer
        p = foot.paragraphs[0] if foot.paragraphs else foot.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{NAME} - Curriculum Vitae - page "); r.font.size = Pt(8.5)
        fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
        rr = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"; rr.append(t); fld.append(rr)
        p._p.append(fld)

    # ---- sections -------------------------------------------------------------------------------------------
    def cv_education(self, entries=None):
        """approved_claims.yaml cv.education[] {institution, dates, degree, loc, details[]}; falls back to settings.yaml education."""
        self._flush_jobs()
        entries = entries if entries is not None else _ac.CV_EDUCATION
        if not entries:
            return self.education()
        self._heading("Education")
        for e in entries:
            p = self._tabbed(space_before=3)
            r = p.add_run(str(e.get("institution", ""))); r.bold = True; r.font.size = Pt(10.5)
            p.add_run("\t" + str(e.get("dates", ""))).font.size = Pt(10)
            p2 = self.doc.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
            p2.paragraph_format.tab_stops.add_tab_stop(Twips(RIGHT_TAB), WD_TAB_ALIGNMENT.RIGHT)
            ri = p2.add_run(str(e.get("degree", ""))); ri.italic = True
            if e.get("loc"):
                ri2 = p2.add_run("\t" + str(e.get("loc"))); ri2.italic = True
            for d in (e.get("details") or []):
                self._bullet(str(d))
        return self

    def experience_section(self, section, entries=None, title=None):
        """One dated-entry section: section in research|professional|field|teaching (title from SECTION_TITLES unless
        given). Entries default to approved_claims experience[] with that `section` (field -> cv.field_experience,
        teaching -> cv.teaching dict items). Each entry: {org|title, dates, role, loc, bullets}."""
        self._flush_jobs()
        if entries is None:
            if section == "field":
                entries = _ac.CV_FIELD
            elif section == "teaching":
                entries = [t for t in _ac.CV_TEACHING if isinstance(t, dict)]
            else:
                entries = _ac.EXPERIENCE_BY_SECTION.get(section, [])
        if not entries:
            return self
        self._heading(title or self.SECTION_TITLES.get(section, section.title()))
        rows = []
        for e in entries:
            bl = e.get("bullets", {}) or {}
            bullets = list(bl.values()) if isinstance(bl, dict) else [str(x) for x in bl]
            rows.append((str(e.get("org") or e.get("title") or ""), self._norm_dates(str(e.get("dates", ""))),
                         str(e.get("role") or e.get("title") or ""), str(e.get("loc", "")), bullets))
        for j in sorted(rows, key=lambda j: self._date_key(j[1]), reverse=True):
            self._render_job(*j)
        return self

    def _hanging(self, text, number=None, size=10):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Pt(18); p.paragraph_format.first_line_indent = Pt(-18)
        if number is not None:
            p.add_run(f"{number}.\t").font.size = Pt(size)
            p.paragraph_format.tab_stops.add_tab_stop(Pt(18))
        p.add_run(text).font.size = Pt(size)
        return p

    def list_section(self, title, items, numbered=False):
        """A section of plain items (publications, presentations, certifications, grants/awards, service).
        Items print verbatim (hanging indent); numbered for publications. Empty list = section omitted."""
        self._flush_jobs()
        items = [str(x) for x in (items or []) if str(x).strip()]
        if not items:
            return self
        self._heading(title)
        for i, t in enumerate(items, 1):
            self._hanging(t, number=i if numbered else None)
        return self

    def publications(self, items=None):
        return self.list_section("Publications", _ac.CV_PUBLICATIONS if items is None else items, numbered=True)

    def presentations(self, items=None):
        return self.list_section("Presentations", _ac.CV_PRESENTATIONS if items is None else items)

    def certifications(self, items=None):
        return self.list_section("Certifications & Training", _ac.CV_CERTIFICATIONS if items is None else items)

    def technical_skills(self, lines=None):
        """Skills lines: cv.technical_skills (or skills_lines) as (label, text) pairs."""
        self._flush_jobs()
        lines = lines if lines is not None else list(_ac.CV_SKILLS.values())
        if not lines:
            return self
        self._heading("Technical Skills")
        for label, val in lines:
            p = self.doc.add_paragraph(); p.paragraph_format.space_after = Pt(1)
            p.add_run(label + " : ").bold = True
            p.add_run(val)
        return self

    def teaching(self):
        """cv.teaching: dict items become dated entries, string items a plain list."""
        dicts = [t for t in _ac.CV_TEACHING if isinstance(t, dict)]
        strs = [str(t) for t in _ac.CV_TEACHING if not isinstance(t, dict)]
        if dicts:
            self.experience_section("teaching", dicts)
            if strs:
                for t in strs: self._hanging(t)
        elif strs:
            self.list_section("Teaching & Mentoring", strs)
        return self

    def grants_awards(self, items=None):
        return self.list_section("Grants, Fellowships & Awards", _ac.CV_GRANTS if items is None else items)

    def service(self, items=None):
        return self.list_section("Service & Outreach", _ac.CV_SERVICE if items is None else items)

    def references(self, text=None):
        self._flush_jobs()
        text = _ac.CV_REFERENCES if text is None else text
        if not text:
            return self
        self._heading("References")
        for t in (text if isinstance(text, list) else [text]):
            self._hanging(str(t))
        return self

    def standard(self, summary=None):
        """The default section order for a research / PhD application. Any section with nothing approved is skipped.
        Call the section methods yourself for a different order."""
        if summary:
            self.summary(summary)
        self.cv_education()
        self.experience_section("research")
        self.experience_section("professional")
        self.publications()
        self.presentations()
        self.experience_section("field")
        self.technical_skills()
        self.certifications()
        self.teaching()
        self.grants_awards()
        self.service()
        self.references()
        return self

    def save(self, path):
        self._flush_jobs()
        self._set_core_properties()
        self.doc.core_properties.title = f"{NAME} - Curriculum Vitae"
        self.doc.core_properties.subject = "jobkit:cv"
        self.doc.save(path)
        print("wrote", path)
        return path

