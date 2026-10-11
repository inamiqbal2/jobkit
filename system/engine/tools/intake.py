#!/usr/bin/env python3
"""INTAKE - read the documents the person dropped into "My Job Search/intake/" (v0.4).

The person puts old resumes / CVs, cover letters, essays, a thesis, papers, reports, performance reviews,
certificates into the intake folder and says "read my documents". This tool:
  1. lists the folder (every file; hidden files, README.md and the _extracted/ folder are ignored)
  2. extracts plain text from .docx (python-docx, incl. tables), .pdf (pypdf), .txt and .md (as is)
  3. writes intake/_extracted/<name>.txt  (one per readable file; re-extracted only when the source is newer)
     and intake/_extracted/INDEX.md     (file, type, size, words, status - the thing Claude reads first)
  4. NEVER modifies, moves or deletes an original. Nothing leaves the computer.
Claude then reads the extracted text and DRAFTS evidence files / a skills inventory / a publications list, every
item marked DRAFT with a `Source: intake/<file>` line, and confirms them in the onboarding interview
(system/engine/skills/onboarding/SKILL.md, "Confirmation mode"). Nothing from intake is an approved claim until the person
has heard it read back and said yes.

  <jobkit python> intake.py            # extract what is new or changed, write the index, print the summary
  <jobkit python> intake.py --list     # only list the folder (no extraction)
  <jobkit python> intake.py --force    # re-extract everything
Exit 0 always (a file that cannot be read is listed as such; it never stops the others)."""
import sys as _sys0
try: _sys0.stdout.reconfigure(encoding="utf-8", errors="replace")   # Windows: never die on a non-cp1252 character
except Exception: pass
import os, sys, re, datetime, json

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from userconfig import INTAKE, INTAKE_EXTRACTED, WS

SUPPORTED = {".docx": "Word", ".pdf": "PDF", ".txt": "text", ".md": "Markdown"}
SKIP_NAMES = {"README.md", ".gitkeep", ".DS_Store", "Thumbs.db", "INDEX.md"}
INDEX = os.path.join(INTAKE_EXTRACTED, "INDEX.md")
INDEX_JSON = os.path.join(INTAKE_EXTRACTED, "index.json")


def _docx(path):
    from docx import Document
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    # headers/footers often carry the name and contact line
    for s in d.sections:
        for hf in (s.header, s.footer):
            try:
                parts += [p.text for p in hf.paragraphs if p.text.strip()]
            except Exception:
                pass
    return "\n".join(parts)


def _pdf(path):
    from pypdf import PdfReader
    r = PdfReader(path)
    out = []
    for i, page in enumerate(r.pages):
        try:
            out.append(page.extract_text() or "")
        except Exception as e:
            out.append(f"[page {i + 1}: could not extract text: {type(e).__name__}]")
    return "\n\n".join(out)


def _plain(path):
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return open(path, encoding=enc).read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    return open(path, encoding="utf-8", errors="replace").read()


def extract(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        return _docx(path)
    if ext == ".pdf":
        return _pdf(path)
    return _plain(path)


def _clean(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip() + "\n"


def listing():
    """Every candidate file in intake/ (top level and sub-folders, except _extracted/)."""
    out = []
    if not os.path.isdir(INTAKE):
        return out
    for root, dirs, files in os.walk(INTAKE):
        dirs[:] = sorted(d for d in dirs if not d.startswith(".") and d != "_extracted")
        for f in sorted(files):
            if f in SKIP_NAMES or f.startswith(".") or f.startswith("~$"):
                continue
            out.append(os.path.join(root, f))
    return out


def run(force=False, list_only=False):
    files = listing()
    os.makedirs(INTAKE_EXTRACTED, exist_ok=True)
    rows = []
    for src in files:
        rel = os.path.relpath(src, INTAKE)
        ext = os.path.splitext(src)[1].lower()
        size = os.path.getsize(src)
        kind = SUPPORTED.get(ext)
        if not kind:
            rows.append(dict(file=rel, type=ext or "(none)", size=size, words=0, extracted="", status="skipped: not a .docx/.pdf/.txt/.md file"))
            continue
        if list_only:
            rows.append(dict(file=rel, type=kind, size=size, words=0, extracted="", status="listed"))
            continue
        dst = os.path.join(INTAKE_EXTRACTED, re.sub(r"[^\w.-]+", "_", rel.replace(os.sep, "__")) + ".txt")
        fresh = os.path.exists(dst) and os.path.getmtime(dst) >= os.path.getmtime(src)
        if fresh and not force:
            txt = open(dst, encoding="utf-8").read()
            rows.append(dict(file=rel, type=kind, size=size, words=len(txt.split()), extracted=os.path.relpath(dst, WS), status="ok (unchanged)"))
            continue
        try:
            txt = _clean(extract(src))
            with open(dst, "w", encoding="utf-8") as fh:
                fh.write(f"# extracted from intake/{rel} on {datetime.date.today().isoformat()} - DRAFT MATERIAL, not approved claims\n\n")
                fh.write(txt)
            words = len(txt.split())
            status = "ok" if words else "ok but EMPTY (a scanned PDF? no text layer - ask the person for a Word or text version)"
            rows.append(dict(file=rel, type=kind, size=size, words=words, extracted=os.path.relpath(dst, WS), status=status))
        except Exception as e:
            rows.append(dict(file=rel, type=kind, size=size, words=0, extracted="", status=f"could not read: {type(e).__name__}: {str(e)[:80]}"))
    if not list_only:
        with open(INDEX, "w", encoding="utf-8") as fh:
            fh.write(f"# intake index - {datetime.date.today().isoformat()} (written by intake.py; the originals are never changed)\n\n")
            fh.write("Every extracted file is DRAFT material. Facts become usable only after the read-back in the interview.\n\n")
            fh.write("| file | type | size | words | extracted text | status |\n|---|---|---|---|---|---|\n")
            for r in rows:
                fh.write(f"| {r['file']} | {r['type']} | {r['size']:,} B | {r['words']:,} | {r['extracted']} | {r['status']} |\n")
        with open(INDEX_JSON, "w", encoding="utf-8") as fh:
            json.dump(rows, fh, indent=1)
    return rows


def main(argv):
    force = "--force" in argv; list_only = "--list" in argv
    if not os.path.isdir(INTAKE):
        os.makedirs(INTAKE_EXTRACTED, exist_ok=True)
        print(f"intake folder created (it was missing): {INTAKE}\nNothing to read yet - drop documents there and say 'read my documents'.")
        return 0
    rows = run(force=force, list_only=list_only)
    if not rows:
        print(f"intake folder is empty: {INTAKE}\nDrop resumes, CVs, letters, essays, papers, reports or reviews there (.docx .pdf .txt .md) and say 'read my documents'.")
        return 0
    ok = [r for r in rows if r["status"].startswith("ok")]
    print(f"intake: {len(rows)} file(s) in {INTAKE}" + ("" if list_only else f"; {len(ok)} readable, {sum(r['words'] for r in ok):,} words extracted"))
    for r in rows:
        print(f"  {r['file']:<50} {r['type']:<9} {r['words']:>7,} words  {r['status']}" + (f"  -> {r['extracted']}" if r['extracted'] else ""))
    if not list_only:
        print(f"index: {os.path.relpath(INDEX, WS)}  (read every extracted file in full before drafting)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
